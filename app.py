import re
import pandas as pd
import streamlit as st
from pypdf import PdfReader
from docx import Document
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity

# ---------------------------------------------------------
# Page setup and visual theme
# ---------------------------------------------------------
st.set_page_config(
    page_title="AI  | Resume Screener",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
:root { color-scheme: light; }
.block-container { padding-top: 1.7rem; padding-bottom: 3rem; max-width: 1400px; }
[data-testid="stAppViewContainer"] { background: #f6f8fc; }
[data-testid="stHeader"] { background: rgba(246,248,252,0.9); }
.hero {
    padding: 1.8rem 2rem; border-radius: 20px; color: white;
    background: linear-gradient(120deg, #14213d 0%, #233f73 60%, #4169a8 100%);
    margin-bottom: 1.25rem;
}
.hero-kicker { color: #c9d9ff; font-size: .78rem; letter-spacing: .14em; text-transform: uppercase; font-weight: 700; }
.hero h1 { color: white; margin: .35rem 0 .45rem 0; font-size: 2.25rem; }
.hero p { color: #e0eaff; margin: 0; font-size: 1rem; }
.section-card {
    background: white; border: 1px solid #e4e9f2; border-radius: 16px;
    padding: 1.1rem 1.25rem; margin-bottom: 1rem;
}
.metric-label { color: #68758b; font-size: .82rem; margin-bottom: .3rem; }
.metric-value { color: #172b4d; font-size: 1.65rem; font-weight: 750; }
.metric-note { color: #758198; font-size: .76rem; margin-top: .2rem; }
div[data-testid="stMetric"] {
    background: white; border: 1px solid #e4e9f2; padding: 1rem 1.1rem;
    border-radius: 14px;
}
div[data-testid="stMetricLabel"] { color: #63718a; }
div[data-testid="stMetricValue"] { color: #172b4d; }
.stButton > button[kind="primary"] {
    border-radius: 10px; border: 0; min-height: 2.8rem; font-weight: 650;
}
.stDownloadButton > button { border-radius: 10px; }
div[data-testid="stExpander"] { border: 1px solid #e2e8f0; border-radius: 12px; background: white; }
.small-muted { color: #6b778c; font-size: .85rem; }
.pill {
    display: inline-block; padding: .28rem .58rem; margin: .15rem .18rem .15rem 0;
    border-radius: 999px; background: #eaf2ff; color: #214f91; font-size: .8rem;
}
.pill-missing { background: #fff2e8; color: #9a4a0a; }
hr { border-color: #e4e9f2; }
</style>
""", unsafe_allow_html=True)

if "screening_data" not in st.session_state:
    st.session_state["screening_data"] = None

# ---------------------------------------------------------
# Models and resume processing
# ---------------------------------------------------------
@st.cache_resource
def load_model():
    return SentenceTransformer("all-MiniLM-L6-v2")

def extract_text(file):
    file.seek(0)
    filename = file.name.lower()
    if filename.endswith(".pdf"):
        reader = PdfReader(file)
        text = " ".join(page.extract_text() or "" for page in reader.pages)
    elif filename.endswith(".docx"):
        document = Document(file)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells if cell.text.strip())
        text = " ".join(parts)
    else:
        text = ""
    return text.strip()

def extract_skills(text):
    skills_list = [
        "python", "sql", "machine learning", "deep learning", "pandas", "numpy",
        "scikit-learn", "tensorflow", "pytorch", "excel", "power bi", "tableau",
        "nlp", "data analysis", "data visualization", "java", "c++", "communication",
        "statistics", "data cleaning", "data engineering", "artificial intelligence",
        "natural language processing", "computer vision", "streamlit", "matplotlib",
        "seaborn", "git", "github", "aws", "azure", "docker", "mongodb", "mysql",
        "eda", "problem-solving", "problem solving", "machine learning", "reporting",
        "dashboard", "dashboards"
    ]
    text = text.lower()
    found = []
    for skill in skills_list:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text) and skill not in found:
            found.append(skill)
    return found

def generate_explanation(job_description, resume_text):
    required_skills = extract_skills(job_description)
    candidate_skills = extract_skills(resume_text)
    matched = [s for s in required_skills if s in candidate_skills]
    missing = [s for s in required_skills if s not in candidate_skills]
    if not required_skills:
        return (
            "No skills from the app's predefined skill list were detected in the job description. "
            "Add specific required skills to improve skill matching. Review the resume manually."
        )
    ratio = len(matched) / len(required_skills) * 100
    if ratio >= 75:
        fit_note = "The resume contains many of the listed skills and may merit closer review."
    elif ratio >= 40:
        fit_note = "The resume contains some listed skills; compare the candidate's project evidence with the role requirements."
    else:
        fit_note = "Few listed skills were detected; check whether relevant skills are described using different wording."
    matched_text = ", ".join(matched) if matched else "No listed skills detected"
    missing_text = ", ".join(missing) if missing else "No listed skills were missing from the text"
    return (
        f"{fit_note} Skills found in both the job description and resume: {matched_text}. "
        f"Skills not detected in the resume text (may still be possessed by the candidate): {missing_text}. "
        "Absence from the document does not prove the candidate lacks a skill; verify claims and experience manually."
    )

# ---------------------------------------------------------
# Header and sidebar
# ---------------------------------------------------------
st.markdown("""
<div class="hero">
<div class="hero-kicker">AI-POWERED RECRUITMENT</div>
<h1>AI Resume Screener</h1>
<p>AI-powered resume ranking, skill matching, and candidate insights.</p>
</div>
""", unsafe_allow_html=True)


with st.sidebar:
    st.markdown("## 📄 AI Resume screener")
    st.caption("Resume screening workspace")
    st.markdown("---")
    st.markdown("### How it works")
    st.markdown("1. Add the job description\n2. Upload PDF or DOCX resumes\n3. Run the comparison\n4. Review evidence and export results")
    st.markdown("---")
    st.caption("Privacy note: uploaded resume content is processed by this app and its model libraries. Use test resumes while developing.")

# ---------------------------------------------------------
# Inputs
# ---------------------------------------------------------
st.markdown("### 1 · Define the role")
st.markdown('<p class="small-muted">Include responsibilities and specific skills to improve the quality of the comparison.</p>', unsafe_allow_html=True)
job_description = st.text_area(
    "Job description",
    height=180,
    placeholder="Example: Junior Data Analyst with Python, SQL, Pandas, Excel, data cleaning, EDA, statistics, and data visualization...",
    label_visibility="collapsed",
)

st.markdown("### 2 · Add candidate resumes")
uploaded_files = st.file_uploader(
    "Upload resumes (PDF or DOCX; you can select multiple files)",
    type=["pdf", "docx"],
    accept_multiple_files=True,
    help="Text-based PDFs and DOCX files are supported. Scanned PDFs may need OCR before screening.",
)

left, right = st.columns([1, 3])
with left:
    screen_clicked = st.button("🔎 Screen resumes", type="primary", use_container_width=True)
with right:
    st.markdown('<p class="small-muted">Tip: use the same job description for every candidate in a comparison.</p>', unsafe_allow_html=True)

if screen_clicked:
    if not job_description.strip():
        st.warning("Please enter a job description.")
    elif not uploaded_files:
        st.warning("Please upload at least one resume.")
    else:
        resume_texts, valid_files = [], []
        for file in uploaded_files:
            try:
                text = extract_text(file)
                if not text:
                    st.warning(f"No readable text found in {file.name}. Scanned PDFs may require OCR.")
                    continue
                resume_texts.append(text)
                valid_files.append(file.name)
            except Exception as error:
                st.warning(f"Could not read {file.name}: {error}")

        if not resume_texts:
            st.error("No readable resumes were found.")
        else:
            try:
                with st.spinner("Loading the semantic model and comparing resumes…"):
                    similarity_model = load_model()
                    job_embedding = similarity_model.encode(job_description, normalize_embeddings=True)
                    resume_embeddings = similarity_model.encode(resume_texts, normalize_embeddings=True)
                    similarities = cosine_similarity([job_embedding], resume_embeddings)[0]

                required_skills = extract_skills(job_description)
                results = []
                explanations = {}
                progress = st.progress(0, text="Preparing candidate reports…")
                for i, resume_text in enumerate(resume_texts):
                    candidate_skills = extract_skills(resume_text)
                    matched = [s for s in required_skills if s in candidate_skills]
                    missing = [s for s in required_skills if s not in candidate_skills]
                    explanations[valid_files[i]] = generate_explanation(job_description, resume_text)
                    results.append({
                        "Resume": valid_files[i],
                        "AI Similarity (%)": round(float(similarities[i]) * 100, 2),
                        "Matched Skills": ", ".join(matched) or "None found",
                        "Missing Skills": ", ".join(missing) or "None from the listed skills",
                        "AI Explanation": explanations[valid_files[i]],
                    })
                    progress.progress((i + 1) / len(resume_texts), text=f"Processed {i + 1} of {len(resume_texts)} resumes")
                progress.empty()

                results_df = pd.DataFrame(results).sort_values(
                    by="AI Similarity (%)", ascending=False
                ).reset_index(drop=True)
                results_df.insert(0, "Rank", range(1, len(results_df) + 1))
                st.session_state["screening_data"] = {
                    "results_df": results_df,
                    "resume_texts": resume_texts,
                    "valid_files": valid_files,
                    "job_description": job_description,
                    "explanations": explanations,
                }
                st.success(f"Screening complete. {len(results_df)} resume(s) processed.")
            except Exception as error:
                st.error(f"Could not screen resumes: {error}")

# ---------------------------------------------------------
# Results dashboard
# ---------------------------------------------------------
if st.session_state["screening_data"] is not None:
    data = st.session_state["screening_data"]
    results_df = data["results_df"]
    total = len(results_df)
    average = results_df["AI Similarity (%)"].mean() if total else 0
    top_score = results_df["AI Similarity (%)"].max() if total else 0
    skill_count = len(extract_skills(data["job_description"]))

    st.markdown("---")
    st.markdown("### 3 · Screening overview")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Resumes reviewed", total)
    m2.metric("Top similarity", f"{top_score:.2f}%")
    m3.metric("Average similarity", f"{average:.2f}%")
    m4.metric("Skills identified", skill_count)

    st.markdown("### 4 · Candidate ranking")
    st.caption("Ranked by semantic text similarity. Scores are not hiring probabilities or a final suitability decision.")
    display_df = results_df[["Rank", "Resume", "AI Similarity (%)", "Matched Skills", "Missing Skills"]]
    st.dataframe(
        display_df,
        use_container_width=True,
        hide_index=True,
        column_config={
            "Rank": st.column_config.NumberColumn("Rank", format="%d"),
            "AI Similarity (%)": st.column_config.ProgressColumn(
                "Semantic similarity", min_value=0, max_value=100, format="%.2f%%"
            ),
            "Matched Skills": st.column_config.TextColumn("Matched skills", width="large"),
            "Missing Skills": st.column_config.TextColumn("Not found in text", width="large"),
        },
    )

    chart_col, info_col = st.columns([1.7, 1])
    
    
    
    with chart_col:
        st.markdown("### Similarity comparison")

        chart_data = results_df[["Resume", "AI Similarity (%)"]].copy()
        chart_data["AI Similarity (%)"] = pd.to_numeric(
            chart_data["AI Similarity (%)"], errors="coerce"
        )
        chart_data = chart_data.dropna(
            subset=["AI Similarity (%)"]
        ).sort_values("AI Similarity (%)", ascending=False)

        if not chart_data.empty:
            chart_data["Candidate"] = [
                f"Candidate {i}"
                for i in range(1, len(chart_data) + 1)
            ]
            chart_data = chart_data.set_index("Candidate")

            st.bar_chart(
                chart_data[["AI Similarity (%)"]],
                horizontal=True,
                height=max(250, len(chart_data) * 65),
                use_container_width=True,
            )
        else:
            st.info("No valid similarity scores available.")



    with info_col:
        st.markdown("### Reading the results")
        st.markdown(
            "- **Similarity** compares overall text meaning.\n"
            "- **Matched skills** are predefined keywords found in both texts.\n"
            "- **Not found in text** means the keyword was not detected; it does not prove the candidate lacks it."
        )
        st.download_button(
            "⬇️ Download screening report (CSV)",
            data=results_df.to_csv(index=False).encode("utf-8"),
            file_name="resume_screening_results.csv",
            mime="text/csv",
            use_container_width=True,
        )

    st.markdown("### 5 · Candidate review")
    st.caption("Review the original resume and verify evidence before making any decision.")
    for _, row in results_df.iterrows():
        with st.expander(f"Rank {row['Rank']}  ·  {row['Resume']}  ·  {row['AI Similarity (%)']:.2f}%"):
            c1, c2 = st.columns(2)
            with c1:
                st.markdown("**Skills detected in both**")
                if row["Matched Skills"] == "None found":
                    st.info("No predefined skill keywords matched.")
                else:
                    st.markdown(" ".join(
                        f'<span class="pill">{s.strip()}</span>'
                        for s in row["Matched Skills"].split(",")
                    ), unsafe_allow_html=True)
            with c2:
                st.markdown("**Job skills not found in resume text**")
                if row["Missing Skills"] in ("None from the listed skills", ""):
                    st.success("No listed skill keywords were missing.")
                else:
                    st.markdown(" ".join(
                        f'<span class="pill pill-missing">{s.strip()}</span>'
                        for s in row["Missing Skills"].split(",")
                    ), unsafe_allow_html=True)
            st.markdown("**Review summary**")
            st.write(row["AI Explanation"])

    st.markdown("---")
    st.caption(
        "Responsible use: this tool assists comparison only. Semantic similarity and keyword matching can be incomplete or biased. "
        "Do not use the score as an automatic hiring decision; review original materials consistently and fairly."
    )
else:
    st.markdown("""
    <div class="section-card">
      <strong>Your results will appear here</strong>
      <p class="small-muted">Enter a role description, upload candidate resumes, and select <em>Screen resumes</em> to build the comparison dashboard.</p>
    </div>
    """, unsafe_allow_html=True)
